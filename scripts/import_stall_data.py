#!/usr/bin/env python3
"""Rebuild scripts/seed/stalls.json from the original records in sources/.

Firestore is the source of truth once seeded; this is only for a fresh
database. Amounts, partners and order come from Project Vita Information.docx
(transcribed into DOC_STALLS and DOC_MENUS), the rest from Stall Results.xlsx.
sources/ is not committed.
"""

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

SOURCE = ROOT / "sources" / "Stall Results.xlsx"

OUTPUT = ROOT / "scripts" / "seed" / "stalls.json"

# Page 1 of the information doc, completed stalls in chronological order, then
# the upcoming one. `slug` matches the workbook row where there is one; `date`
# overrides the workbook's date.

DOC_STALLS = [
    {"slug": "sports-day-2023", "amountRaised": 44450, "partner": "RGH + Inventure"},
    {"slug": "eat-play-love-2023", "amountRaised": 45000, "partner": "Mitti Cafe"},
    {"slug": "inmun-2024", "amountRaised": 47000, "partner": "Care For All"},
    {"slug": "inmun-2025", "amountRaised": 227317, "partner": "Care For All"},
    {"slug": "sports-day-2025", "amountRaised": 72480, "partner": "Care For All"},
    {
        "slug": "hack-club-daydream-2025",
        "title": "Hack Club Daydream 2025",
        "date": "2025",
        "amountRaised": 234384,
        "partner": "Hack Club",
    },
    {"slug": "production-2025", "amountRaised": 42150, "partner": "RGH BA Program"},
    # The doc lists the partner as "TBD".
    {"slug": "swim-meet-2026", "amountRaised": 18370, "partner": None},
    {
        "slug": "codeday-2026",
        "status": "upcoming",
        "date": "2026",
        "amountRaised": 61500,
        "partner": "CodeDay",
    },
]

# Menus shown in the doc, kept only where the workbook lists the same item: the
# doc supplies the name and price, the workbook the category and count. Entries
# are (doc name, doc price, workbook item); a None price falls back to the
# workbook's. Stalls missing here have no menu in the doc and use the workbook's.

DOC_MENUS = {
    # Doc heading: "Winter Carnival 2023" (The Pit Stop).
    "eat-play-love-2023": [
        ("Soya Sticks", 50, "Soya sticks"),
        ("Chips and Dip", 70, "Lay's chips (4 flavours)"),
        ("Nachos and Dip", 100, "Nachos with salsa, cheese & chipotle dips"),
        ("French Fries", 100, "French fries"),
        ("Chicken Nuggets", 150, "Chicken nuggets"),
        ("Gatorade", 100, "Gatorade"),
    ],
    "inmun-2024": [
        ("Chip n Dip", 70, "Lay's with dip (20 g plate)"),
        ("Nachos and Salsa", 100, "Nachos with dip (50 g plate)"),
        ("Fries", 150, "Fries (100 g plate)"),
        ("Veg Momo", 120, "Veg momos (5 pc)"),
        ("Chicken Nuggets", 150, "Chicken nuggets (5 pc)"),
        ("Chicken Momo", 150, "Chicken momos (5 pc)"),
        ("Gatorade", 50, "Gatorade (170 ml glass)"),
        ("Blue Lagoon", 100, "Blue Lagoon (240 ml)"),
        ("Brownie", 70, "Brownie"),
    ],
    "inmun-2025": [
        ("Lays Chips", 70, "Lay's (variety, 30 g)"),
        ("Nachos and Salsa", 100, "Nachos (variety, 30 g)"),
        ("French Fries", 100, "Fries (50 g)"),
        ("Cheesy Fries", 120, "Cheese fries (50 g)"),
        ("Butter Popcorn", 50, "Butter popcorn (30 g)"),
        ("Cheese Popcorn", 70, "Cheese popcorn (30 g)"),
        ("Vegetable Momos", 120, "Veg momos (4 pc)"),
        ("Chicken Momos", 150, "Chicken momos (4 pc)"),
        ("Chicken Nuggets", 150, "Chicken nuggets (4 pc)"),
        ("Orange Juice", 30, "Orange juice (200 ml)"),
        ("Lime Soda", 30, "Lime soda (200 ml)"),
        ("Gatorade Cup", 50, "Gatorade (200 ml)"),
        ("Iced Tea", 70, "Lemon iced tea (200 ml)"),
        ("Blue Lagoon", 70, "Blue Lagoon (200 ml)"),
        ("Chocolate Brownie", 100, "Chocolate brownie"),
        ("Craft your own Sundae", 100, "Ice cream sundae base"),
        ("Brownie add-on: ice cream", 50, "Extra ice cream"),
        ("Tier 1 topping: chocolate sauce", 10, "Topping tier 1: chocolate sauce"),
        ("Tier 1 topping: caramel sauce", 10, "Topping tier 1: caramel sauce"),
        ("Tier 1 topping: strawberry sauce", 10, "Topping tier 1: strawberry sauce"),
        ("Tier 2 topping: M&Ms", 20, "Topping tier 2: M&Ms"),
        ("Tier 2 topping: sprinkles", 20, "Topping tier 2: sprinkles"),
        ("Tier 2 topping: chocolate chips", 20, "Topping tier 2: chocolate chips"),
        ("Tier 3 topping: gummy worms", 30, "Topping tier 3: gummy worms"),
        ("Tier 3 topping: gummy strips", 30, "Topping tier 3: gummy strips"),
        ("Tier 3 topping: gummy bites", 30, "Topping tier 3: gummy bites"),
        ("Tier 4 topping: Oreo crumble", 50, "Topping tier 4: Oreo crumble"),
        ("Tier 4 topping: KitKat", 50, "Topping tier 4: KitKat crumble"),
    ],
    "production-2025": [
        ("Butter Popcorn", 100, "Butter Popcorn"),
        ("Cheese Popcorn", 120, "Cheese Popcorn"),
        ("Nachos & Salsa", 150, "Nachos n Salsa"),
        ("Veg Momo", 170, "Veg Momos"),
        ("Chicken Momo", 200, "Chicken Momos"),
        ("Fresh Lemonade", 100, "Lemonade"),
        ("Fresh Orange Juice", 100, "Orange Juice"),
    ],
    # The doc lists the food without prices, so the workbook's bill prices stay.
    "codeday-2026": [
        ("Lunch: veg rice bowls", None, "Lunch"),
        ("Snack: popcorn", None, "Popcorn"),
        ("Dinner: veg burger", None, "Dinner"),
    ],
}

UNKNOWN_PARTNER = {"name": "Partner not recorded", "recorded": False}

PHOTO_SETS = {
    "swim-meet-2026": {
        "heroImage": "/stalls/swim-meet-2026/hero.webp",
        "heroImageAlt": "Students serving customers at the Project Vita Swim Meet stall",
        "gallery": [
            {
                "src": "/stalls/swim-meet-2026/hero.webp",
                "alt": "Students serving customers at the Project Vita Swim Meet stall",
                "crop": "center 42%",
            },
            {
                "src": "/stalls/swim-meet-2026/cooking.webp",
                "alt": "A student preparing food at the Swim Meet stall",
                "crop": "center",
            },
            {
                "src": "/stalls/swim-meet-2026/plate.webp",
                "alt": "Food being plated at the Swim Meet stall",
                "crop": "center",
            },
            {
                "src": "/stalls/swim-meet-2026/counter.webp",
                "alt": "Students working at the Swim Meet stall counter",
                "crop": "center",
            },
            *(
                {
                    "src": f"/stalls/swim-meet-2026/gallery/img-{number}.webp",
                    "alt": alt,
                    "crop": "center",
                }
                for number, alt in [
                    ("3431", "A volunteer logging orders on a laptop before the stall opens"),
                    ("3439", "Popcorn kernels popping in the stall's popcorn maker"),
                    ("3441", "A volunteer adding kernels to the popcorn maker"),
                    ("3443", "The Swim Meet stall seen from the walkway, with volunteers setting up"),
                    ("3444", "Students crowding around the stall's order table"),
                    ("3445", "A queue forming under the blue canopy at the Swim Meet stall"),
                    ("3447", "Customers gathered in front of the stall under the canopy"),
                    ("3448", "Volunteers working behind the stall tables"),
                    ("3449", "A volunteer seasoning a tray of fries"),
                    ("3450", "A volunteer shaking seasoning over fries"),
                    ("3451", "Fries being plated at the Swim Meet stall"),
                    ("3453", "Ketchup being added to a plate of fries"),
                    ("3455", "The order and payment table with its QR code"),
                    ("3456", "A volunteer holding a finished plate of fries"),
                    ("3457", "A volunteer in a hairnet tending the popcorn maker"),
                    ("3458", "A smiling volunteer at the Swim Meet stall"),
                    ("3460", "Customers paying at the order table"),
                    ("3462", "The stall's canopy, cooler and supplies"),
                    ("3463", "Popcorn being tossed with seasoning in a tub"),
                    ("3464", "A volunteer in a hairnet seasoning popcorn"),
                    ("3466", "Popcorn being scooped into a serving cup"),
                    ("3469", "A plate of nachos being served"),
                    ("3470", "A volunteer drizzling sauce over nachos"),
                ]
            ),
        ],
    }
}


def rows_as_dicts(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]


def slugify(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def normalise_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    match = re.search(r"\b(20\d{2})\b", str(value or ""))
    return match.group(1) if match else "Date not recorded"


def number_or_none(value):
    return value if isinstance(value, (int, float)) else None


def format_rupees(amount):
    """Indian digit grouping (₹2,27,317), matching the site's formatter."""
    digits = str(round(amount))
    head, tail = digits[:-3], digits[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return "₹" + ",".join(groups + [tail])


def build_story(title, amount_raised, ngo, status):
    raised = format_rupees(amount_raised)
    if status == "upcoming":
        return f"{title} is the next stall, expected to raise {raised} with {ngo['name']}."
    if ngo["recorded"]:
        return f"{title} raised {raised} in support of {ngo['name']}."
    return f"{title} raised {raised}. Its partner organisation is still to be confirmed."


def build_partner(name):
    partner_name = str(name or "").strip()
    if not partner_name:
        return dict(UNKNOWN_PARTNER)

    # Descriptions live in the partners collection (scripts/seed/partners.json).
    return {"name": partner_name, "recorded": True}


def build_menu(slug, workbook_menu):
    if slug not in DOC_MENUS:
        return workbook_menu

    items_by_name = {item["name"]: item for item in workbook_menu}
    menu = []
    for doc_name, doc_price, workbook_name in DOC_MENUS[slug]:
        item = dict(items_by_name[workbook_name])
        item["name"] = doc_name
        if doc_price is not None:
            item["price"] = doc_price
        menu.append(item)
    return menu


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    stall_rows = rows_as_dicts(workbook["Stalls"])
    menu_rows = rows_as_dicts(workbook["Menus"])

    menus = defaultdict(list)
    for row in menu_rows:
        stall_number = number_or_none(row.get("No."))
        item_name = row.get("Item")
        if stall_number is None or not item_name:
            continue
        menus[int(stall_number)].append(
            {
                "name": str(item_name).strip(),
                "category": str(row.get("Category") or "Other").strip(),
                "price": number_or_none(row.get("Price (₹)")),
                "quantitySold": number_or_none(row.get("Qty sold")),
            }
        )

    workbook_rows = {}
    for row in stall_rows:
        workbook_number = number_or_none(row.get("No."))
        title = row.get("Stall")
        if workbook_number is not None and title:
            workbook_rows[slugify(str(title))] = (int(workbook_number), str(title).strip(), row)

    stalls = []
    for stall_number, record in enumerate(DOC_STALLS, start=1):
        slug = record["slug"]
        workbook_number, title, row = workbook_rows.pop(slug, (None, record.get("title"), {}))
        date = record.get("date") or normalise_date(row.get("Date"))
        ngo = build_partner(record["partner"])
        status = record.get("status", "completed")
        menu = build_menu(slug, menus[workbook_number])
        plates_served = number_or_none(row.get("Items served"))
        if status == "upcoming":
            # Nothing has been sold yet; the workbook's counts are a plan.
            menu = [{**item, "quantitySold": None} for item in menu]
            plates_served = None

        stall = {
            "slug": slug,
            "title": title,
            "status": status,
            "published": True,
            "pinned": False,
            "sortOrder": stall_number,
            "date": date,
            "story": build_story(title, record["amountRaised"], ngo, status),
            "ngo": ngo,
            "impact": {
                "amountRaised": record["amountRaised"],
                "expenses": number_or_none(row.get("Total expenses (₹)")),
                "platesServed": plates_served,
                "itemsBasis": None if status == "upcoming" else row.get("Items basis"),
            },
            "menu": menu,
            "gallery": [],
        }
        stall.update(PHOTO_SETS.get(slug, {}))
        stalls.append(stall)

    if workbook_rows:
        raise SystemExit(f"Workbook stalls missing from DOC_STALLS: {', '.join(workbook_rows)}")

    stalls.sort(key=lambda stall: stall["sortOrder"], reverse=True)
    OUTPUT.write_text(json.dumps(stalls, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(stalls)} stalls and {sum(len(stall['menu']) for stall in stalls)} menu items")
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
